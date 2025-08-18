#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import sys
import json
from docx import Document
import os

def extract_docx_content(file_path):
    """
    从DOCX文件中提取文本内容
    """
    try:
        # 打开DOCX文档
        doc = Document(file_path)
        
        content = []
        
        # 提取段落文本
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                content.append(text)
        
        # 提取表格文本
        for table in doc.tables:
            for row in table.rows:
                row_content = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_content.append(cell_text)
                if row_content:
                    content.append(" | ".join(row_content))
        
        # 获取文件名作为标题
        title = os.path.splitext(os.path.basename(file_path))[0]
        
        # 合并所有内容
        full_content = "\n".join(content)
        
        return {
            "title": title,
            "content": full_content,
            "source": file_path,
            "type": "docx"
        }
        
    except Exception as e:
        return {
            "error": f"处理DOCX文件失败: {str(e)}"
        }

def main():
    if len(sys.argv) != 2:
        print(json.dumps({"error": "请提供DOCX文件路径"}))
        sys.exit(1)
    
    file_path = sys.argv[1]
    
    if not os.path.exists(file_path):
        print(json.dumps({"error": f"文件不存在: {file_path}"}))
        sys.exit(1)
    
    result = extract_docx_content(file_path)
    print(json.dumps(result, ensure_ascii=False))

if __name__ == "__main__":
    main() 