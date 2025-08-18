#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel转PPT转换脚本
用于将网页内容总结的Excel文件转换为PPT格式
"""

import sys
import os
import json
from datetime import datetime

try:
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
except ImportError as e:
    print(f"❌ 缺少必要的Python库: {e}")
    print("请安装: pip3 install openpyxl python-pptx")
    sys.exit(1)

def create_ppt_from_excel(excel_file, ppt_file):
    """
    从Excel文件创建PPT
    """
    try:
        # 加载Excel文件
        print(f"📖 正在读取Excel文件: {excel_file}")
        wb = load_workbook(excel_file)
        ws = wb.active
        
        # 创建PPT
        print("🎨 正在创建PPT...")
        prs = Presentation()
        
        # 添加标题页
        title_slide_layout = prs.slide_layouts[0]  # 标题页布局
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = "网页内容总结"
        subtitle.text = "AI智能总结报告"
        
        # 设置标题样式
        title_para = title.text_frame.paragraphs[0]
        title_para.font.size = Pt(44)
        title_para.font.bold = True
        title_para.font.color.rgb = RGBColor(46, 134, 171)  # #2E86AB
        
        # 设置副标题样式
        subtitle_para = subtitle.text_frame.paragraphs[0]
        subtitle_para.font.size = Pt(24)
        subtitle_para.font.color.rgb = RGBColor(128, 128, 128)
        
        # 添加内容页
        content_slide_layout = prs.slide_layouts[1]  # 标题和内容布局
        
        # 读取Excel内容并添加到PPT
        current_section = None
        current_content = []
        
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            if not any(cell for cell in row if cell):
                continue
                
            cell_value = str(row[0]) if row[0] else ""
            
            # 检查是否是新的部分标题
            if is_section_title(cell_value):
                # 保存前一个部分
                if current_section and current_content:
                    add_content_slide(prs, content_slide_layout, current_section, current_content)
                
                # 开始新部分
                current_section = cell_value
                current_content = []
            else:
                # 添加到当前部分
                if cell_value.strip():
                    current_content.append(cell_value.strip())
        
        # 保存最后一个部分
        if current_section and current_content:
            add_content_slide(prs, content_slide_layout, current_section, current_content)
        
        # 如果没有内容，添加一个默认页
        if len(prs.slides) == 1:
            slide = prs.slides.add_slide(content_slide_layout)
            title = slide.shapes.title
            content = slide.placeholders[1]
            
            title.text = "主要内容"
            content.text = "网页内容总结已完成，请查看详细内容。"
        
        # 保存PPT
        print(f"💾 正在保存PPT文件: {ppt_file}")
        prs.save(ppt_file)
        
        print(f"✅ PPT文件已成功生成: {ppt_file}")
        return True
        
    except Exception as e:
        print(f"❌ 转换失败: {e}")
        return False

def is_section_title(text):
    """
    判断是否是部分标题
    """
    if not text:
        return False
    
    # 检查是否以数字和点开头
    if len(text) > 2 and text[1] == '.' and text[0].isdigit():
        return True
    
    # 检查是否包含关键词
    keywords = ["概述", "要点", "信息", "结论", "总结", "主要内容", "关键要点", "重要信息"]
    for keyword in keywords:
        if keyword in text:
            return True
    
    return False

def add_content_slide(prs, layout, title_text, content_list):
    """
    添加内容页
    """
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    # 设置标题
    title.text = title_text
    
    # 设置标题样式
    title_para = title.text_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(46, 134, 171)
    
    # 设置内容
    content_text = ""
    for i, item in enumerate(content_list, 1):
        if item.startswith('•'):
            content_text += f"{item}\n"
        elif item[0].isdigit() or item[0].isalpha():
            content_text += f"• {item}\n"
        else:
            content_text += f"{item}\n"
    
    content.text = content_text
    
    # 设置内容样式
    content_para = content.text_frame.paragraphs[0]
    content_para.font.size = Pt(18)
    content_para.font.color.rgb = RGBColor(0, 0, 0)

def main():
    """
    主函数
    """
    print("=== Excel转PPT转换工具 ===")
    print()
    
    if len(sys.argv) != 3:
        print("使用方法: python3 convert_to_ppt.py <excel_file> <ppt_file>")
        print("示例: python3 convert_to_ppt.py input.xlsx output.pptx")
        sys.exit(1)
    
    excel_file = sys.argv[1]
    ppt_file = sys.argv[2]
    
    # 检查输入文件是否存在
    if not os.path.exists(excel_file):
        print(f"❌ 错误: Excel文件不存在: {excel_file}")
        sys.exit(1)
    
    # 检查输出目录
    output_dir = os.path.dirname(ppt_file)
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    # 执行转换
    if create_ppt_from_excel(excel_file, ppt_file):
        print("🎉 转换完成!")
        sys.exit(0)
    else:
        print("💥 转换失败!")
        sys.exit(1)

if __name__ == "__main__":
    main() 